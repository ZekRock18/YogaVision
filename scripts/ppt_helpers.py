"""YogaVision presentation builder.

Generates a 17-slide professional deck from the real project artifacts in
outputs/ (training curves, confusion matrix, evaluation metrics).

Run:  python scripts/make_presentation.py
"""
from __future__ import annotations

import os
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

PROJECT_ROOT = Path(__file__).resolve().parent.parent

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def rgb(r: int, g: int, b: int) -> RGBColor:
    return RGBColor(r, g, b)


C_BG = rgb(8, 12, 28)
C_NAVY = rgb(10, 20, 60)
C_CARD = rgb(14, 18, 38)
C_BLUE = rgb(37, 99, 235)
C_CYAN = rgb(6, 182, 212)
C_GREEN = rgb(52, 211, 153)
C_WHITE = rgb(255, 255, 255)
C_GREY = rgb(148, 163, 184)
C_DGREY = rgb(71, 85, 105)
C_AMBER = rgb(251, 191, 36)
C_RED = rgb(248, 113, 113)
C_PURPLE = rgb(167, 139, 250)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def add_slide() -> object:
    return prs.slides.add_slide(BLANK)


def bg(slide, color: RGBColor = C_BG) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, l, t, w, h, fill_rgb=None, line_rgb=None, line_width_pt=0.0):
    shape = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.line.fill.background()
    if fill_rgb is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_rgb
    else:
        shape.fill.background()
    if line_rgb is not None:
        shape.line.color.rgb = line_rgb
        shape.line.width = Pt(line_width_pt)
    return shape


def txt(slide, text, l, t, w=None, h=None, size=18, bold=False, color=C_WHITE,
        align=PP_ALIGN.LEFT, italic=False, line_spacing=1.15):
    w = 12 if w is None else w
    h = 1 if h is None else h
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    _write_paragraphs(tf, text.split("\n"), size, bold, color, align, italic, line_spacing)
    return tb


def _write_paragraphs(tf, lines, size, bold, color, align, italic, line_spacing):
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        pPr = p._p.get_or_add_pPr()
        ln_spc = etree.SubElement(pPr, qn("a:lnSpc"))
        spc_pct = etree.SubElement(ln_spc, qn("a:spcPct"))
        spc_pct.set("val", str(int(line_spacing * 100000)))
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color


def badge(slide, text, l, t, w, h, fill_rgb, size=14, text_color=C_WHITE):
    rect(slide, l, t, w, h, fill_rgb)
    txt(slide, text, l, t + h * 0.18, w, h * 0.7, size=size, bold=True,
        color=text_color, align=PP_ALIGN.CENTER)


def card(slide, l, t, w, h, fill_rgb=C_CARD, border_rgb=C_BLUE, border_pt=1.0):
    return rect(slide, l, t, w, h, fill_rgb, border_rgb, border_pt)


def topbar(slide, h=0.06, color=C_CYAN):
    rect(slide, 0, 0, 13.333, h, fill_rgb=color)


def bottombar(slide, h=0.22, color=C_CYAN):
    rect(slide, 0, 7.5 - h, 13.333, h, fill_rgb=color)
    txt(slide, "YogaVision  |  Computer Vision Project", 0.3, 7.28, 12.5, 0.22,
        size=9, color=C_WHITE)


def heading(slide, title, subtitle=None):
    topbar(slide, 0.06, C_CYAN)
    txt(slide, title, 0.4, 0.12, 12.5, 0.6, size=28, bold=True, color=C_WHITE)
    if subtitle:
        txt(slide, subtitle, 0.4, 0.65, 12.5, 0.35, size=13, color=C_GREY, italic=True)


def kv_table(slide, rows, left, top, width, key_w, value_w, row_h=0.5,
             key_color=C_GREEN, value_color=C_WHITE, key_size=12, value_size=12,
             border=C_GREEN, value_align=PP_ALIGN.RIGHT, zebra=True):
    for i, (k, v) in enumerate(rows):
        y = top + i * row_h
        fill = C_CARD if (i % 2 == 0 or not zebra) else rgb(18, 22, 42)
        rect(slide, left, y, width, row_h - 0.05, fill, border, 0.4)
        txt(slide, k, left + 0.18, y + 0.07, key_w, row_h - 0.2,
            size=key_size, bold=True, color=key_color)
        txt(slide, v, left + key_w + 0.25, y + 0.05, value_w, row_h - 0.2,
            size=value_size, bold=True, color=value_color, align=value_align)


def picture_or_placeholder(slide, rel_path, l, t, w, h, label):
    path = PROJECT_ROOT / rel_path
    if path.exists():
        slide.shapes.add_picture(str(path), Inches(l), Inches(t), Inches(w), Inches(h))
        return True
    card(slide, l, t, w, h, C_CARD, C_DGREY, 0.5)
    txt(slide, label, l, t + h / 2 - 0.25, w, 0.5, size=14, color=C_GREY, align=PP_ALIGN.CENTER)
    return False


def metric_cards(slide, items, left, top, width, height, gap=0.25, cols=3,
                 border=C_GREEN, big_size=26, label_size=12):
    card_w = (width - gap * (cols - 1)) / cols
    for i, (big, label, sub) in enumerate(items):
        x = left + (i % cols) * (card_w + gap)
        y = top + (i // cols) * (height + gap)
        card(slide, x, y, card_w, height, C_CARD, border, 1.5)
        txt(slide, big, x + 0.1, y + 0.15, card_w - 0.2, height * 0.4,
            size=big_size, bold=True, color=border, align=PP_ALIGN.CENTER)
        txt(slide, label, x + 0.1, y + height * 0.55, card_w - 0.2, 0.35,
            size=label_size, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        if sub:
            txt(slide, sub, x + 0.1, y + height * 0.78, card_w - 0.2, 0.4,
                size=9.5, color=C_GREY, align=PP_ALIGN.CENTER)


def feature_grid(slide, features, left=0.3, top=1.1, col_w=4.35, row_h=3.0,
                 cols=3, bullets_per_card=6):
    for idx, (color, title, bullets) in enumerate(features):
        x = left + (idx % cols) * col_w
        y = top + (idx // cols) * row_h
        card(slide, x, y, col_w - 0.25, row_h - 0.15, C_CARD, color, 1.2)
        rect(slide, x, y, col_w - 0.25, 0.55, color)
        txt(slide, title, x + 0.15, y + 0.1, col_w - 0.55, 0.38, size=13, bold=True,
            color=C_WHITE, align=PP_ALIGN.CENTER)
        for i, b in enumerate(bullets[:bullets_per_card]):
            prefix = "+ " if i < 3 else "  "
            txt(slide, prefix + b, x + 0.2, y + 0.68 + i * 0.36, col_w - 0.6, 0.32,
                size=10, color=C_WHITE)


def info_grid(slide, items, left=0.3, top=1.1, col_w=6.6, row_h=2.1, cols=2):
    for i, (title, desc, color) in enumerate(items):
        x = left + (i % cols) * col_w
        y = top + (i // cols) * row_h
        card(slide, x, y, col_w - 0.3, row_h - 0.15, C_CARD, color, 1.0)
        rect(slide, x, y, 0.55, row_h - 0.15, color)
        txt(slide, title, x + 0.7, y + 0.1, col_w - 1.0, 0.42, size=13, bold=True, color=color)
        txt(slide, desc, x + 0.7, y + 0.52, col_w - 1.05, row_h - 0.75, size=10.5, color=C_GREY)


def save(output_rel: str) -> Path:
    out = PROJECT_ROOT / output_rel
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out
